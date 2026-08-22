#!/usr/bin/env python3
import argparse
import csv
import json
import re
from pathlib import Path

ALIASES = {
    "manufacturer": ["manufacturer", "vendor", "brand", "厂家", "制造商", "品牌", "厂商"],
    "part_number": ["partnumber", "part_number", "part number", "model", "type", "型号", "订货号", "物料号", "零件号", "料号"],
    "quantity": ["qty", "quantity", "count", "数量", "个数"],
    "device_tag": ["device_tag", "devicetag", "dt", "tag", "device", "设备代号", "设备标识", "位号", "设备编号"],
    "variant": ["variant", "变体"],
    "page": ["page", "页", "页号"],
    "function_location": ["function_location", "functionlocation", "location", "安装位置", "功能位置"],
    "description": ["description", "name", "名称", "描述", "品名"],
}


def key(value):
    value = "" if value is None else str(value).strip().lower()
    return re.sub(r"[\s_\-:/\\（）()\[\]{}]+", "", value)


def text(value):
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def qty(value):
    s = text(value)
    if not s:
        return 1
    try:
        n = int(float(s))
        return max(n, 1)
    except Exception:
        return 1


def detect_columns(headers):
    normalized = {key(h): h for h in headers if h is not None}
    out = {}
    for field, aliases in ALIASES.items():
        for alias in aliases:
            k = key(alias)
            if k in normalized:
                out[field] = normalized[k]
                break
    return out


def read_csv(path):
    last = None
    for enc in ("utf-8-sig", "utf-8", "gb18030"):
        try:
            with path.open("r", encoding=enc, newline="") as f:
                return list(csv.DictReader(f))
        except UnicodeDecodeError as e:
            last = e
    raise last


def read_xlsx(path, sheet=None):
    try:
        from openpyxl import load_workbook
    except ImportError:
        raise SystemExit("XLSX requires openpyxl. Run: python -m pip install openpyxl")
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb[sheet] if sheet else wb.active
    rows = ws.iter_rows(values_only=True)
    headers = [text(x) for x in next(rows, [])]
    result = []
    for values in rows:
        result.append({headers[i]: values[i] if i < len(values) else None for i in range(len(headers))})
    return result


def main():
    ap = argparse.ArgumentParser(description="Normalize BOM CSV/XLSX into an EPLAN on-demand parts queue")
    ap.add_argument("input")
    ap.add_argument("--sheet", default=None)
    ap.add_argument("--out", default=None)
    args = ap.parse_args()

    path = Path(args.input)
    if not path.exists():
        raise SystemExit(f"Input not found: {path}")

    suffix = path.suffix.lower()
    if suffix == ".csv":
        raw = read_csv(path)
    elif suffix in (".xlsx", ".xlsm"):
        raw = read_xlsx(path, args.sheet)
    else:
        raise SystemExit("Supported BOM formats: .csv, .xlsx, .xlsm")

    headers = list(raw[0].keys()) if raw else []
    cols = detect_columns(headers)
    if "part_number" not in cols:
        raise SystemExit(f"Could not detect a part/model column. Headers: {headers}")

    rows = []
    unique = {}
    for i, src in enumerate(raw, start=2):
        manufacturer = text(src.get(cols.get("manufacturer", ""), ""))
        part_number = text(src.get(cols["part_number"], ""))
        if not part_number:
            continue
        row = {
            "row": i,
            "manufacturer": manufacturer,
            "part_number": part_number,
            "quantity": qty(src.get(cols.get("quantity", ""), 1)),
            "device_tag": text(src.get(cols.get("device_tag", ""), "")),
            "variant": text(src.get(cols.get("variant", ""), "")) or "1",
            "page": text(src.get(cols.get("page", ""), "")),
            "function_location": text(src.get(cols.get("function_location", ""), "")),
            "description": text(src.get(cols.get("description", ""), "")),
        }
        row["status"] = "ready" if manufacturer else "missing_manufacturer"
        rows.append(row)
        ukey = (key(manufacturer), part_number.strip().lower())
        if ukey not in unique:
            unique[ukey] = {
                "manufacturer": manufacturer,
                "part_number": part_number,
                "rows": [],
                "status": "resolve" if manufacturer else "missing_manufacturer",
            }
        unique[ukey]["rows"].append(i)

    doc = {
        "source": str(path.resolve()),
        "sheet": args.sheet,
        "detected_columns": cols,
        "row_count": len(rows),
        "unique_part_count": len(unique),
        "rows": rows,
        "parts": list(unique.values()),
    }

    payload = json.dumps(doc, ensure_ascii=False, indent=2)
    if args.out:
        out = Path(args.out)
        out.parent.mkdir(parents=True, exist_ok=True)
        out.write_text(payload, encoding="utf-8")
        print(out)
    else:
        print(payload)


if __name__ == "__main__":
    main()
